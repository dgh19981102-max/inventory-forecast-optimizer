"""Excel in / Excel out.

Input workbook (see examples/input_template.xlsx)
    Settings       parameter | value | description
    Products       one row per SKU: costs, pack sizes, lead time, stock on hand, open orders
    DemandHistory  week_start | <sku 1> | <sku 2> | ...   (weekly units, oldest first)

Output workbook
    Summary          headline numbers + cost comparison of the three ways to order
    Order Plan       SKU x week order quantities (week 1 = place this week)
    Plan Detail      per SKU-week: forecast P10/P50/P90, order, arrival, projected stock
    Forecast Accuracy backtest table that decided which forecasting model is used
    Assumptions      the settings used
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRODUCT_COLUMNS = {
    "sku": "Product code",
    "description": "Product name",
    "sell_price": "Selling price per unit",
    "unit_cost": "Purchase cost per unit",
    "space_per_unit": "Storage space per unit (1 = one standard slot)",
    "holding_cost_week": "Cost of holding 1 unit for 1 week (capital + storage)",
    "lost_sale_penalty": "Cost of 1 unit of unmet demand (lost margin + goodwill)",
    "case_size": "Supplier pack size (orders are whole packs)",
    "moq": "Minimum order quantity per order (units)",
    "lead_time_weeks": "Weeks between placing an order and receiving it",
    "fixed_order_cost": "Fixed cost per order line (admin, delivery)",
    "on_hand": "Units in stock now",
    "on_order_week1": "Open orders arriving at start of plan week 1",
    "on_order_week2": "Open orders arriving at start of plan week 2",
}

SETTINGS_DESCRIPTIONS = {
    "warehouse_capacity_space": "Warehouse capacity in storage-space units (soft limit: overflow is rented)",
    "weekly_budget": "Maximum purchase spend per week",
    "overflow_cost_per_space_week": "Cost per space unit per week stored above capacity",
    "planning_horizon_weeks": "Weeks to plan ahead (2-8)",
    "service_level_target": "Target service level for the safety-stock method (e.g. 0.95)",
    "random_seed": "Random seed for the cost comparison",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOW_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")


@dataclass
class PlanInput:
    settings: dict
    products: pd.DataFrame          # index = sku (str)
    history: pd.DataFrame           # index = week_start (Timestamp), columns = sku


# --------------------------------------------------------------------------- reading
def read_input(path: str | Path) -> PlanInput:
    xl = pd.ExcelFile(path)
    s = pd.read_excel(xl, "Settings")
    settings = {str(k).strip(): v for k, v in zip(s["parameter"], s["value"])}
    for k in ("planning_horizon_weeks", "random_seed"):
        settings[k] = int(settings[k])
    for k in ("warehouse_capacity_space", "weekly_budget", "overflow_cost_per_space_week", "service_level_target"):
        settings[k] = float(settings[k])

    prod = pd.read_excel(xl, "Products", dtype={"sku": str})
    missing = set(PRODUCT_COLUMNS) - set(prod.columns)
    if missing:
        raise ValueError(f"Products sheet is missing columns: {sorted(missing)}")
    prod["sku"] = prod["sku"].astype(str).str.strip()
    prod = prod.set_index("sku")
    for c in PRODUCT_COLUMNS:
        if c not in ("sku", "description"):
            prod[c] = pd.to_numeric(prod[c]).fillna(0)
    for c in ("case_size", "moq", "lead_time_weeks"):
        prod[c] = prod[c].astype(int)
    if (prod["lead_time_weeks"] > 2).any() or (prod["lead_time_weeks"] < 1).any():
        raise ValueError("lead_time_weeks must be 1 or 2 in this template (add on_order_week3.. columns for longer)")

    hist = pd.read_excel(xl, "DemandHistory")
    hist = hist.rename(columns={hist.columns[0]: "week_start"})
    hist["week_start"] = pd.to_datetime(hist["week_start"])
    hist = hist.set_index("week_start").sort_index()
    hist.columns = [str(c).strip() for c in hist.columns]
    absent = [s for s in prod.index if s not in hist.columns]
    if absent:
        raise ValueError(f"No demand history for: {absent}")
    hist = hist[list(prod.index)].fillna(0).astype(float)
    if len(hist) < 60:
        raise ValueError("Need at least 60 weeks of history (52 for seasonality + backtest)")
    return PlanInput(settings, prod, hist)


# --------------------------------------------------------------------------- writing helpers
def _style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autowidth(ws, min_w=8, max_w=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, width + 2))


def _write_df(ws, df: pd.DataFrame, start_row=1, index=True, number_formats: dict | None = None):
    d = df.reset_index() if index else df
    for j, c in enumerate(d.columns, 1):
        ws.cell(row=start_row, column=j, value=str(c))
    for i, row in enumerate(d.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating, np.integer)):
                v = v.item()
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if number_formats and d.columns[j - 1] in number_formats:
                cell.number_format = number_formats[d.columns[j - 1]]
            elif isinstance(v, float):
                cell.number_format = "#,##0.00"
            elif hasattr(v, "year"):
                cell.number_format = "yyyy-mm-dd"
    _style_header(ws, start_row, len(d.columns))


# --------------------------------------------------------------------------- template
def write_template(path: str | Path, settings: dict, products: pd.DataFrame, history: pd.DataFrame):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [
        "Replenishment planner - input template",
        "",
        "1. Settings: warehouse capacity, weekly budget, planning horizon ... (yellow cells).",
        "2. Products: one row per product. Costs, pack size, minimum order, lead time, stock on hand, open orders.",
        "3. DemandHistory: weekly units sold per product, oldest week first, at least 60 weeks (ideally 2 years).",
        "   First column = week start date; one column per product code (same codes as in Products).",
        "",
        "Then run:  python scripts/plan_from_excel.py <this file> <output file>",
        "or upload it in the web app:  streamlit run app/streamlit_app.py",
        "",
        "Column definitions (Products):",
    ] + [f"   {k}: {v}" for k, v in PRODUCT_COLUMNS.items()]
    for i, line in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Settings")
    sdf = pd.DataFrame({"parameter": list(SETTINGS_DESCRIPTIONS),
                        "value": [settings[k] for k in SETTINGS_DESCRIPTIONS],
                        "description": list(SETTINGS_DESCRIPTIONS.values())})
    _write_df(ws, sdf, index=False)
    for r in range(2, len(sdf) + 2):
        ws.cell(row=r, column=2).fill = INPUT_FILL
    _autowidth(ws)

    ws = wb.create_sheet("Products")
    p = products.reset_index()[list(PRODUCT_COLUMNS)]
    _write_df(ws, p, index=False)
    for r in range(2, len(p) + 2):
        for c in range(3, len(PRODUCT_COLUMNS) + 1):
            ws.cell(row=r, column=c).fill = INPUT_FILL
    _autowidth(ws)
    ws.freeze_panes = "C2"

    ws = wb.create_sheet("DemandHistory")
    h = history.copy()
    h.index = h.index.rename("week_start")
    _write_df(ws, h, index=True, number_formats={c: "0" for c in h.columns})
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 12
    wb.save(path)


# --------------------------------------------------------------------------- output
def write_output(path: str | Path, res: dict):
    """`res` is the dict returned by planner.make_plan."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Replenishment plan"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("Plan starts (week 1)", res["weeks"][0].date().isoformat()),
        ("Planning horizon (weeks)", len(res["weeks"])),
        ("Products", len(res["products"])),
        ("Forecast model (best in backtest)", res["model"]),
        ("Backtest WAPE of that model", f"{res['model_wape']:.1%}"),
        ("Order lines to place THIS week", int((res["orders"][0] > 0).sum())),
        ("Purchase value of this week's orders", round(float((res["orders"][0] * res["products"]["unit_cost"]).sum()), 2)),
        ("Optimiser solve time (s)", round(res["solve_seconds"], 1)),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    r0 = len(info) + 5
    ws.cell(row=r0 - 1, column=1, value="Expected cost over the planning horizon - same demand scenarios for all three plans").font = Font(bold=True, size=12)
    _write_df(ws, res["comparison"], start_row=r0, index=True)
    note_row = r0 + len(res["comparison"]) + 2
    notes = [
        "How to read: each plan is the full order schedule for the horizon, scored on "
        f"{res['n_eval_scenarios']} demand scenarios built from this data's own past forecast errors.",
        "Rule of thumb = order up to (lead time + 2) weeks of last-4-week average sales, whole packs, MOQ, budget cap.",
        "Costs depend on the cost assumptions in Products/Settings. Re-plan every week with fresh stock levels.",
    ]
    for i, n_ in enumerate(notes):
        ws.cell(row=note_row + i, column=1, value=n_).font = Font(italic=True, color="595959")
    ws.column_dimensions["A"].width = 44
    for c in "BCDE":
        ws.column_dimensions[c].width = 26

    # Order plan
    ws = wb.create_sheet("Order Plan")
    P = res["products"]
    plan = pd.DataFrame(res["orders"].T, index=P.index, columns=[f"Week {k+1}\n{w.date()}" for k, w in enumerate(res["weeks"])])
    plan.insert(0, "description", P["description"])
    plan.insert(1, "lead_time_weeks", P["lead_time_weeks"])
    plan["total_units"] = plan.iloc[:, 2:].sum(axis=1)
    plan["total_cost"] = (plan["total_units"] * P["unit_cost"]).round(2)
    plan = plan.rename_axis("sku")
    _write_df(ws, plan, index=True, number_formats={c: "#,##0" for c in plan.columns if c.startswith("Week") or c == "total_units"})
    for r in range(2, len(plan) + 2):
        ws.cell(row=r, column=4).fill = NOW_FILL
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    _autowidth(ws, max_w=36)

    ws = wb.create_sheet("Plan Detail")
    _write_df(ws, res["detail"], index=False, number_formats={"week_start": "yyyy-mm-dd"})
    ws.freeze_panes = "C2"
    _autowidth(ws)

    ws = wb.create_sheet("Forecast Accuracy")
    _write_df(ws, res["accuracy"].round(3), index=True)
    ws.cell(row=len(res["accuracy"]) + 3, column=1,
            value="WAPE = sum|forecast-actual| / sum(actual); Bias > 0 means over-forecasting. Rolling-origin backtest, 1-6 weeks ahead.").font = Font(italic=True)
    _autowidth(ws)

    ws = wb.create_sheet("Assumptions")
    _write_df(ws, pd.Series(res["settings"], name="value").rename_axis("setting").to_frame(), index=True)
    _autowidth(ws)
    wb.save(path)
